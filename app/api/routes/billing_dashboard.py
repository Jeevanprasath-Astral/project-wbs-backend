"""Group Billing Dashboard — planned vs actual billing (₹ lakhs) rolled up
Group → Company → Division → Project → Milestone.

  GET /billing-dashboard/data    — JSON tree + flat entry list
  GET /billing-dashboard/export  — same data as .xlsx (Summary + Billing Entries)

Source data is ProjectBilling (Financial Settings → Billing):
  Planned amount  = planned_billing_amount
  Actual amount   = actual_billing_amount (only once billed)
  Planned date    = the linked milestone's planned_end
  Actual date     = actual_billing_date

Timeliness of each billing entry:
  Earlier     billed before the planned date (beyond the grace window)
  On time     billed within ±grace_days of the planned date (default 0 = same day)
  Delayed     billed after the planned date (beyond the grace window)
  Upcoming    not billed yet, planned date today or later
  Overdue     not billed yet, planned date already passed
  Unscheduled no planned date (milestone not linked or has no planned end)

Filters:
  group_id / company_id / division_id / project_id — org structure
  status       — one or more statuses, comma-separated
  start_date / end_date — period, on the billing date (actual date if billed,
                 otherwise planned date). Entries with no date are excluded
                 when a period is set.
  grace_days   — tolerance for "On time"
Projects not yet assigned to a division appear under "Unassigned" unless an
org filter is applied.
"""
from collections import OrderedDict
from datetime import date as date_cls
from io import BytesIO
from typing import Optional

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.core.deps import get_current_user
from app.db.database import get_db
from app.models.models import (Company, CustomMilestone, Division, OrgGroup, Project,
                               ProjectBilling, User)

router = APIRouter(prefix="/billing-dashboard", tags=["Billing Dashboard"])

LAKH = 100_000.0
STATUSES = ["Earlier", "On time", "Delayed", "Upcoming", "Overdue", "Unscheduled"]
UNASSIGNED = "Unassigned"
NO_MILESTONE = "Not linked to a milestone"


# ── Helpers ───────────────────────────────────────────────────────────────────
def _parse_date(v: Optional[str]):
    if not v:
        return None
    try:
        return date_cls.fromisoformat(v)
    except ValueError:
        raise HTTPException(400, f"Invalid date '{v}' — use YYYY-MM-DD")


def _parse_statuses(v: Optional[str]):
    if not v:
        return None
    wanted = {s.strip().lower() for s in v.split(",") if s.strip()}
    known = {s.lower(): s for s in STATUSES}
    bad = wanted - set(known)
    if bad:
        raise HTTPException(400, f"Unknown status: {', '.join(sorted(bad))}. "
                                 f"Use: {', '.join(STATUSES)}")
    return {known[s] for s in wanted}


def billing_status(planned, actual, today, grace_days=0):
    """(status, days_vs_plan). days_vs_plan is always measured against the
    planned date: positive = later than plan, negative = ahead of it. For
    entries not billed yet it is measured from today (Upcoming ≤ 0, Overdue > 0)."""
    if planned is None:
        return "Unscheduled", None
    if actual is None:
        return ("Upcoming" if planned >= today else "Overdue"), (today - planned).days
    delta = (actual - planned).days
    if abs(delta) <= grace_days:
        return "On time", delta
    return ("Earlier", delta) if delta < 0 else ("Delayed", delta)


def _lakhs(v: float) -> float:
    return round((v or 0.0) / LAKH, 2)


def _new_node(level, node_id, name):
    return {"level": level, "id": node_id, "name": name,
            "_planned": 0.0, "_actual": 0.0, "entries": 0,
            "_status": {s: {"count": 0, "_planned": 0.0, "_actual": 0.0} for s in STATUSES},
            "_children": OrderedDict()}


def _add(node, e):
    node["_planned"] += e["_planned"]
    node["_actual"] += e["_actual"]
    node["entries"] += 1
    st = node["_status"][e["status"]]
    st["count"] += 1
    st["_planned"] += e["_planned"]
    st["_actual"] += e["_actual"]


def _finish(node):
    """Convert running rupee totals into the public lakh figures."""
    planned, actual = node.pop("_planned"), node.pop("_actual")
    node["planned"] = _lakhs(planned)
    node["actual"] = _lakhs(actual)
    node["variance"] = _lakhs(actual - planned)
    node["achieved_pct"] = round(actual / planned * 100, 1) if planned else None
    node["status"] = {s: {"count": v["count"], "planned": _lakhs(v["_planned"]),
                          "actual": _lakhs(v["_actual"])}
                      for s, v in node.pop("_status").items()}
    children = node.pop("_children")
    if node["level"] != "milestone":
        node["children"] = [_finish(c) for c in children.values()]
    return node


# ── Core ──────────────────────────────────────────────────────────────────────
def build_dashboard(db: Session, user: User, *, group_id=None, company_id=None,
                    division_id=None, project_id=None, status=None, start_date=None,
                    end_date=None, grace_days=0):
    start, end = _parse_date(start_date), _parse_date(end_date)
    if start and end and end < start:
        raise HTTPException(400, "end_date cannot be before start_date")
    if grace_days < 0:
        raise HTTPException(400, "grace_days cannot be negative")
    wanted = _parse_statuses(status)
    today = date_cls.today()

    # ── 1. Projects in scope, with their division/company/group (1 query) ───
    pq = (db.query(Project.id, Project.name, Division.id, Division.name,
                   Company.id, Company.name, OrgGroup.id, OrgGroup.name)
            .outerjoin(Division, Division.id == Project.division_id)
            .outerjoin(Company, Company.id == Division.company_id)
            .outerjoin(OrgGroup, OrgGroup.id == Company.group_id))
    if getattr(user, "is_demo", False):
        pq = pq.filter(Project.is_demo == True)  # noqa: E712
    if group_id:    pq = pq.filter(OrgGroup.id == group_id)
    if company_id:  pq = pq.filter(Company.id == company_id)
    if division_id: pq = pq.filter(Division.id == division_id)
    if project_id:  pq = pq.filter(Project.id == project_id)
    projects = {r[0]: r for r in pq.all()}
    if not projects:
        return _empty(start_date, end_date, grace_days, status)

    # ── 2. Billing entries with their milestone (1 query) ───────────────────
    rows = (db.query(ProjectBilling, CustomMilestone.id, CustomMilestone.name,
                     CustomMilestone.num, CustomMilestone.planned_end)
              .outerjoin(CustomMilestone, CustomMilestone.id == ProjectBilling.milestone_id)
              .filter(ProjectBilling.project_id.in_(list(projects)))
              .all())

    entries = []
    # Status counts for the org/period scope *before* the status filter, so
    # the status chips can still show every bucket while one is selected.
    overview = {s: 0 for s in STATUSES}
    for b, ms_id, ms_name, ms_num, ms_planned_end in rows:
        planned_date = ms_planned_end.date() if ms_planned_end else None
        actual_date = b.actual_billing_date
        st, days = billing_status(planned_date, actual_date, today, grace_days)
        bill_date = actual_date or planned_date
        if start or end:
            if bill_date is None:
                continue
            if start and bill_date < start:
                continue
            if end and bill_date > end:
                continue
        overview[st] += 1
        if wanted and st not in wanted:
            continue
        pid, pname, did, dname, cid, cname, gid, gname = projects[b.project_id]
        planned_amt = float(b.planned_billing_amount or 0)
        actual_amt = float(b.actual_billing_amount) if b.actual_billing_amount is not None else 0.0
        entries.append({
            "id": b.id,
            "group_id": gid, "group": gname or UNASSIGNED,
            "company_id": cid, "company": cname or UNASSIGNED,
            "division_id": did, "division": dname or UNASSIGNED,
            "project_id": pid, "project": pname,
            "milestone_id": ms_id, "milestone": ms_name or NO_MILESTONE, "milestone_num": ms_num,
            "billing_type": b.billing_type,
            "description": b.description,
            "planned_date": planned_date.isoformat() if planned_date else None,
            "actual_date": actual_date.isoformat() if actual_date else None,
            "status": st, "days_variance": days,
            "_planned": planned_amt, "_actual": actual_amt,
            "planned": _lakhs(planned_amt), "actual": _lakhs(actual_amt),
        })

    def _sort_key(e):
        return ((e["group"] == UNASSIGNED, e["group"].lower()),
                (e["company"] == UNASSIGNED, e["company"].lower()),
                (e["division"] == UNASSIGNED, e["division"].lower()),
                e["project"].lower(),
                (e["milestone_id"] is None, e["milestone_num"] or 0, e["milestone"].lower()),
                e["planned_date"] or "9999", e["id"])
    entries.sort(key=_sort_key)

    # ── 3. Roll up ──────────────────────────────────────────────────────────
    total = _new_node("total", None, "All")
    for e in entries:
        _add(total, e)
        g = total["_children"].setdefault(("g", e["group_id"]), _new_node("group", e["group_id"], e["group"]))
        c = g["_children"].setdefault(("c", e["company_id"]), _new_node("company", e["company_id"], e["company"]))
        d = c["_children"].setdefault(("d", e["division_id"]), _new_node("division", e["division_id"], e["division"]))
        p = d["_children"].setdefault(("p", e["project_id"]), _new_node("project", e["project_id"], e["project"]))
        m = p["_children"].setdefault(("m", e["milestone_id"]), _new_node("milestone", e["milestone_id"], e["milestone"]))
        for node in (g, c, d, p, m):
            _add(node, e)

    for e in entries:
        e.pop("_planned"); e.pop("_actual"); e.pop("milestone_num")
    total = _finish(total)
    return {
        "unit": "lakhs",
        "statuses": STATUSES,
        "filters": {"start_date": start_date, "end_date": end_date,
                    "grace_days": grace_days, "status": status},
        "totals": {k: total[k] for k in ("planned", "actual", "variance", "achieved_pct",
                                          "entries", "status")},
        "status_overview": overview,
        "tree": total["children"],
        "entries": entries,
    }


def _empty(start_date, end_date, grace_days, status):
    zero = {s: {"count": 0, "planned": 0.0, "actual": 0.0} for s in STATUSES}
    return {"unit": "lakhs", "statuses": STATUSES,
            "filters": {"start_date": start_date, "end_date": end_date,
                        "grace_days": grace_days, "status": status},
            "totals": {"planned": 0.0, "actual": 0.0, "variance": 0.0, "achieved_pct": None,
                       "entries": 0, "status": zero},
            "status_overview": {s: 0 for s in STATUSES},
            "tree": [], "entries": []}


# ── Endpoints ─────────────────────────────────────────────────────────────────
@router.get("/data")
def dashboard_data(
    group_id: Optional[int] = None,
    company_id: Optional[int] = None,
    division_id: Optional[int] = None,
    project_id: Optional[int] = None,
    status: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    grace_days: int = Query(0, ge=0, le=365),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    return build_dashboard(db, current_user, group_id=group_id, company_id=company_id,
                           division_id=division_id, project_id=project_id, status=status,
                           start_date=start_date, end_date=end_date, grace_days=grace_days)


@router.get("/export")
def dashboard_export(
    group_id: Optional[int] = None,
    company_id: Optional[int] = None,
    division_id: Optional[int] = None,
    project_id: Optional[int] = None,
    status: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    grace_days: int = Query(0, ge=0, le=365),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    data = build_dashboard(db, current_user, group_id=group_id, company_id=company_id,
                           division_id=division_id, project_id=project_id, status=status,
                           start_date=start_date, end_date=end_date, grace_days=grace_days)
    subtitle = _subtitle(db, data, group_id, company_id, division_id, project_id)
    wb = openpyxl.Workbook()
    _summary_sheet(wb.active, data, subtitle)
    _entries_sheet(wb.create_sheet("Billing Entries"), data, subtitle)
    out = BytesIO(); wb.save(out); out.seek(0)
    return StreamingResponse(
        out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=group-billing-dashboard.xlsx"})


# ── Excel ─────────────────────────────────────────────────────────────────────
_HDR_FILL = PatternFill("solid", fgColor="4F46E5")
_TITLE_FILL = PatternFill("solid", fgColor="1F3864")
_LEVEL_FILL = {"group": PatternFill("solid", fgColor="C7D2FE"),
               "company": PatternFill("solid", fgColor="E0E7FF"),
               "division": PatternFill("solid", fgColor="EEF2FF"),
               "project": PatternFill("solid", fgColor="F8FAFC"),
               "milestone": PatternFill("solid", fgColor="FFFFFF")}
_STATUS_FILL = {"Earlier": "DBEAFE", "On time": "DCFCE7", "Delayed": "FEE2E2",
                "Upcoming": "F1F5F9", "Overdue": "FEF3C7", "Unscheduled": "F3F4F6"}
_THIN = Side(style="thin", color="D1D5DB")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_LEVEL_LABEL = {"group": "Group", "company": "Company", "division": "Division",
                "project": "Project", "milestone": "Milestone"}


def _subtitle(db, data, group_id, company_id, division_id, project_id):
    parts = []
    f = data["filters"]
    if f["start_date"] or f["end_date"]:
        parts.append(f"Period: {f['start_date'] or '—'} → {f['end_date'] or '—'}")
    for model, oid, label in ((OrgGroup, group_id, "Group"), (Company, company_id, "Company"),
                              (Division, division_id, "Division"), (Project, project_id, "Project")):
        if oid:
            obj = db.get(model, oid)
            if obj:
                parts.append(f"{label}: {obj.name}")
    if f["status"]:
        parts.append(f"Status: {f['status']}")
    if f["grace_days"]:
        parts.append(f"On time = within ±{f['grace_days']} day(s)")
    parts.append("Amounts in ₹ lakhs")
    return " | ".join(parts)


def _title(ws, title, subtitle, ncols):
    last = get_column_letter(ncols)
    ws.merge_cells(f"A1:{last}1")
    c = ws["A1"]; c.value = title
    c.font = Font(bold=True, color="FFFFFF", size=12); c.fill = _TITLE_FILL
    ws.row_dimensions[1].height = 22
    ws.merge_cells(f"A2:{last}2")
    c = ws["A2"]; c.value = subtitle
    c.font = Font(italic=True, size=9, color="555555")


def _headers(ws, row, headers):
    for col, h in enumerate(headers, 1):
        c = ws.cell(row, col, h)
        c.font = Font(bold=True, color="FFFFFF", size=9); c.fill = _HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _BORDER
    ws.row_dimensions[row].height = 30


def _summary_sheet(ws, data, subtitle):
    ws.title = "Summary"
    headers = ["Level", "Name", "Planned (₹ L)", "Actual (₹ L)", "Variance (₹ L)", "Achieved %",
               "Entries"] + [f"{s} (count)" for s in STATUSES]
    _title(ws, "Group Billing Dashboard — Planned vs Actual", subtitle, len(headers))
    _headers(ws, 4, headers)
    row = 5

    def write(node, depth):
        nonlocal row
        vals = [_LEVEL_LABEL[node["level"]], ("    " * depth) + node["name"], node["planned"],
                node["actual"], node["variance"],
                node["achieved_pct"] if node["achieved_pct"] is not None else "—",
                node["entries"]] + [node["status"][s]["count"] for s in STATUSES]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row, col, v)
            c.border = _BORDER
            c.fill = _LEVEL_FILL[node["level"]]
            c.font = Font(size=9, bold=node["level"] in ("group", "company"))
            if col >= 3:
                c.alignment = Alignment(horizontal="right")
            if col == 5 and isinstance(v, (int, float)) and v < 0:
                c.font = Font(size=9, bold=node["level"] in ("group", "company"), color="B91C1C")
        row += 1
        for ch in node.get("children", []):
            write(ch, depth + 1)

    for g in data["tree"]:
        write(g, 0)
    if not data["tree"]:
        ws.cell(row, 1, "No billing entries for the selected filters").font = Font(italic=True, size=9)
        row += 1
    t = data["totals"]
    vals = ["Total", "", t["planned"], t["actual"], t["variance"],
            t["achieved_pct"] if t["achieved_pct"] is not None else "—", t["entries"]] + \
           [t["status"][s]["count"] for s in STATUSES]
    for col, v in enumerate(vals, 1):
        c = ws.cell(row, col, v)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor="7C3AED"); c.border = _BORDER
        if col >= 3:
            c.alignment = Alignment(horizontal="right")
    widths = [11, 46, 14, 14, 14, 11, 9] + [12] * len(STATUSES)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C5"


def _entries_sheet(ws, data, subtitle):
    headers = ["Group", "Company", "Division", "Project", "Milestone", "Billing Type", "Description",
               "Planned Date", "Actual Date", "Planned (₹ L)", "Actual (₹ L)", "Status",
               "Days vs Plan"]
    _title(ws, "Billing Entries", subtitle, len(headers))
    _headers(ws, 4, headers)
    row = 5
    for e in data["entries"]:
        vals = [e["group"], e["company"], e["division"], e["project"], e["milestone"],
                e["billing_type"] or "—", e["description"] or "", e["planned_date"] or "—",
                e["actual_date"] or "—", e["planned"], e["actual"], e["status"],
                e["days_variance"] if e["days_variance"] is not None else "—"]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row, col, v)
            c.font = Font(size=9); c.border = _BORDER
            if col == 12:
                c.fill = PatternFill("solid", fgColor=_STATUS_FILL[e["status"]])
            if col in (10, 11, 13):
                c.alignment = Alignment(horizontal="right")
        row += 1
    if not data["entries"]:
        ws.cell(row, 1, "No billing entries for the selected filters").font = Font(italic=True, size=9)
    widths = [18, 18, 20, 26, 26, 14, 30, 12, 12, 13, 13, 12, 11]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"
